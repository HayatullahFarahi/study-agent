#!/usr/bin/env python3
"""Turn a learning material into marker-annotated text plus a manifest-ready outline.

Sources: .pdf .pptx .docx .md .txt .html .srt .vtt, audio/video files, .zip archives,
folders of these, and YouTube video or playlist URLs.

Outputs (default folder: <source folder>/_extracted/):
  <name>.txt           text with location markers: "=== PAGE n ===", "=== SLIDE n: title ===",
                       "=== CHAPTER n: title [hh:mm:ss] ===" / "=== PART n [hh:mm:ss] ===", "[hh:mm:ss] ..."
  <name>.outline.json  detected chapters/sections with "loc" (pages/slides/time) and "lines" ranges

Prints a one-line JSON summary on success.
Exit codes: 0 ok, 2 bad input, 3 missing dependency, 4 nothing extracted.
"""

from __future__ import annotations

import argparse
import bisect
import json
import math
import re
import shutil
import sys
import zipfile
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

GENERATOR = "ingest-material/extract.py v1"
TEXT_EXT = {".md", ".markdown", ".txt"}
HTML_EXT = {".html", ".htm"}
SUB_EXT = {".srt", ".vtt"}
MEDIA_EXT = {
    ".mp4",
    ".mkv",
    ".mov",
    ".webm",
    ".avi",
    ".m4v",
    ".mp3",
    ".m4a",
    ".wav",
    ".flac",
    ".ogg",
    ".opus",
    ".aac",
}
SUPPORTED = {".pdf", ".pptx", ".docx"} | TEXT_EXT | HTML_EXT | SUB_EXT | MEDIA_EXT
CONVERT_HINTS = {
    ".ppt": "Re-save it as .pptx in PowerPoint, or run: soffice --headless --convert-to pptx <file>",
    ".doc": "Re-save it as .docx in Word, or run: soffice --headless --convert-to docx <file>",
    ".key": "Export it from Keynote: File > Export To > PowerPoint (or PDF)",
    ".pages": "Export it from Pages: File > Export To > Word (or PDF)",
    ".epub": "Convert it to PDF first, e.g. with Calibre: ebook-convert book.epub book.pdf",
}
YOUTUBE_RE = re.compile(r"^https?://([a-z0-9-]+\.)?(youtube\.com|youtu\.be)/", re.I)
NUMBERING_RE = re.compile(
    r"^\s*(?:(?:chapter|part|section|appendix|lesson|module|unit)\s+(?:\d+(?:\.\d+)*|[ivxlcdm]+|[a-z])\b\s*[:.\-–—]?\s*"
    r"|\d+(?:\.\d+)*[.:)]?(?:\s+|$))",
    re.I,
)
SKIP_TITLE_RE = re.compile(
    r"^(cover|cover page|title page|half title|copyright|copyright page|table of contents|contents|brief contents|"
    r"detailed contents|list of (figures|tables|illustrations)|dedication|acknowledge?ments|index|colophon|"
    r"bibliography|references|endnotes|about the authors?|also by\b.*|praise for\b.*)$",
    re.I,
)
CHAPTER_LINE_RE = re.compile(
    r"^\s*(chapter|part)\s+(\d+|[ivxlcdm]+)\b[\s:.\-–—]*(.*)$", re.I
)
TIMESTAMP = r"(?:\d{1,2}:)?\d{1,2}:\d{2}[.,]\d{1,3}"
CUE_RE = re.compile(rf"^\s*({TIMESTAMP})\s*-->\s*({TIMESTAMP})")
TAG_RE = re.compile(r"<[^>]*>")
LINE_BREAKS_RE = re.compile(r"\r\n|[\r\x0b\x0c\x85\u2028\u2029]")
YDL_BASE = {"quiet": True, "no_warnings": True, "noprogress": True, "noplaylist": True}


class ExtractError(Exception):
    def __init__(self, message: str, code: int = 2):
        super().__init__(message)
        self.code = code


class MissingDependency(Exception):
    def __init__(self, package: str, purpose: str):
        super().__init__(f"{package} (needed for {purpose})")
        self.package = package


def need(module: str, package: str, purpose: str):
    try:
        return __import__(module)
    except ImportError as exc:
        raise MissingDependency(package, purpose) from exc


def progress(message: str) -> None:
    print(message, file=sys.stderr, flush=True)


def hms(seconds: float | None) -> str:
    total = max(0, int(seconds or 0))
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "material"


def safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "-", text).strip("-") or "material"


def pretty(stem: str) -> str:
    text = stem.replace("_", " ")
    if " " not in text:
        text = text.replace("-", " ")
    return " ".join(text.split())


def display_path(path: Path) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return resolved.as_posix()


def natural_key(path: Path) -> list:
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", path.as_posix())
    ]


def strip_numbering(title: str) -> str:
    return NUMBERING_RE.sub("", title, count=1).strip()


def squash(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def roman_to_int(text: str) -> int:
    values = {"i": 1, "v": 5, "x": 10, "l": 50, "c": 100, "d": 500, "m": 1000}
    total = highest = 0
    for char in reversed(text.lower()):
        value = values[char]
        total = total - value if value < highest else total + value
        highest = max(highest, value)
    return total


class Doc:
    """Output lines plus heading marks and location anchors (page / slide / seconds) keyed by line number."""

    def __init__(self, title: str, kind: str, locator: str):
        self.title = title
        self.kind = kind
        self.locator = locator
        self.detected = "none"
        self.duration: float | None = None
        self.meta_id: str | None = None
        self.lines: list[str] = []
        self.marks: list[dict] = []
        self.warnings: list[str] = []
        self._anchor_lines: list[int] = []
        self._anchor_pos: list = []
        self._marker_lines: set[int] = set()

    def add(self, text: str = "") -> int:
        self.lines.extend(
            line.rstrip() for line in LINE_BREAKS_RE.sub("\n", text).split("\n")
        )
        return len(self.lines)

    def add_anchor(self, text: str, pos, marker: bool = True) -> int:
        self.lines.append(LINE_BREAKS_RE.sub(" ", text).replace("\n", " ").rstrip())
        self._anchor_lines.append(len(self.lines))
        self._anchor_pos.append(pos)
        if marker:
            self._marker_lines.add(len(self.lines))
        return len(self.lines)

    def content_end(self, first: int, last: int) -> int:
        """Step back over trailing blank lines and page/slide markers so ranges end on real content."""
        while last > first and (
            not self.lines[last - 1].strip() or last in self._marker_lines
        ):
            last -= 1
        return last

    def mark(self, level: int, title: str, line: int, start=None, end=None) -> None:
        self.marks.append(
            {
                "level": level,
                "title": " ".join(title.split()),
                "line": line,
                "start": start,
                "end": end,
            }
        )

    def warn(self, message: str) -> None:
        if len(self.warnings) < 20:
            self.warnings.append(message)

    def pos_at(self, line: int):
        if not self._anchor_pos:
            return None
        index = bisect.bisect_right(self._anchor_lines, line) - 1
        return self._anchor_pos[max(index, 0)]

    def loc(self, first: int, last: int, mark: dict | None = None) -> str | None:
        if self.locator == "time":
            start = (
                mark["start"]
                if mark and mark.get("start") is not None
                else self.pos_at(first)
            )
            end = (
                mark["end"]
                if mark and mark.get("end") is not None
                else self.pos_at(last)
            )
            return (
                None
                if start is None
                else f"{hms(start)}-{hms(end if end is not None else start)}"
            )
        if self.locator in ("page", "slide"):
            start, end = self.pos_at(first), self.pos_at(last)
            if start is None:
                return None
            one, many = ("p.", "pp.") if self.locator == "page" else ("slide", "slides")
            return f"{one} {start}" if start == end else f"{many} {start}-{end}"
        return None

    def span_loc(self) -> str | None:
        if self.locator == "time":
            return f"00:00:00-{hms(self.duration)}"
        return self.loc(1, len(self.lines))


# ---------------------------------------------------------------- structure


def build_structure(doc: Doc, keep_all: bool = False) -> tuple[list[dict], list[str]]:
    marks = [dict(m) for m in doc.marks if m["title"]]
    if not marks:
        return [], []
    for previous, current in zip(marks, marks[1:]):
        current["line"] = max(current["line"], previous["line"])
    top = min(m["level"] for m in marks)
    marks = [m for m in marks if m["level"] <= top + 1]
    total = len(doc.lines)
    for i, mark in enumerate(marks):
        last = next(
            (m["line"] - 1 for m in marks[i + 1 :] if m["level"] <= mark["level"]),
            total,
        )
        mark["last"] = doc.content_end(mark["line"], max(last, mark["line"]))

    def unit(unit_id: str, mark: dict) -> dict:
        return {
            "id": unit_id,
            "title": mark["title"],
            "loc": doc.loc(mark["line"], mark["last"], mark),
            "lines": [mark["line"], mark["last"]],
            "_mark": mark,
        }

    structure: list[dict] = []
    skipped: list[str] = []
    chapter = None
    skipping = False
    for mark in marks:
        unwanted = not keep_all and bool(
            SKIP_TITLE_RE.match(strip_numbering(mark["title"]))
        )
        if mark["level"] == top:
            skipping = unwanted
            chapter = None
            if unwanted:
                skipped.append(mark["title"])
                continue
        elif skipping:
            continue
        elif unwanted:
            skipped.append(mark["title"])
            continue
        if chapter is None:
            chapter = unit(f"ch-{len(structure) + 1}", mark)
            chapter["sections"] = []
            structure.append(chapter)
        else:
            chapter["sections"].append(
                unit(f"{len(structure)}.{len(chapter['sections']) + 1}", mark)
            )

    for chapter in structure:
        sections = chapter["sections"]
        if sections and sections[0]["lines"][0] > chapter["lines"][0]:
            first = sections[0]
            first["lines"][0] = chapter["lines"][0]
            first["loc"] = doc.loc(first["lines"][0], first["lines"][1], first["_mark"])
        for item in [chapter, *sections]:
            item.pop("_mark", None)
    return structure, skipped


def count_units(structure: list[dict]) -> int:
    return sum(len(chapter.get("sections") or []) or 1 for chapter in structure)


def find_heading_line(
    lines: list[str], title: str, first: int, last: int
) -> int | None:
    want = squash(strip_numbering(title)) or squash(title)
    if len(want) < 4:
        return None
    probe = want[:40]
    for number in range(first, min(last, len(lines)) + 1):
        got = squash(strip_numbering(lines[number - 1]))
        if got and (
            probe in got[: len(probe) + 30] or (len(got) >= 8 and want.startswith(got))
        ):
            above = lines[number - 2].strip() if number > first else ""
            return (
                number - 1 if above and not squash(strip_numbering(above)) else number
            )  # e.g. "Chapter 1" label line
    return None


def add_markdown_marks(doc: Doc) -> None:
    fenced = False
    for number, line in enumerate(doc.lines, 1):
        if re.match(r"^\s*(```|~~~)", line):
            fenced = not fenced
            continue
        match = None if fenced else re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if match:
            doc.mark(len(match.group(1)), match.group(2), number)
    if doc.marks:
        doc.detected = "headings"
        titles = [m["title"] for m in doc.marks if m["level"] == 1]
        if len(titles) == 1:
            doc.title = titles[0]


# ---------------------------------------------------------------- documents


def pdf_title(reader) -> str | None:
    try:
        title = (reader.metadata.title or "").strip() if reader.metadata else ""
    except Exception:
        return None
    if len(title) < 3 or re.search(
        r"^(untitled|microsoft (word|powerpoint) - )|\.(docx?|pdf|indd|tex)$",
        title,
        re.I,
    ):
        return None
    return title


def extract_pdf(path: Path) -> Doc:
    pypdf = need("pypdf", "pypdf", "PDF files")
    try:
        reader = pypdf.PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError(
                "The PDF is password-protected. Remove the password first (e.g. qpdf --decrypt in.pdf out.pdf)."
            )
        pages = list(reader.pages)
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError(f"Could not open the PDF: {exc}") from exc

    doc = Doc(pdf_title(reader) or pretty(path.stem), "pdf", "page")
    try:
        labels = list(reader.page_labels)
    except Exception:
        labels = []
    total = len(pages)
    page_lines: list[int] = []
    chars = 0
    for number, page in enumerate(pages, 1):
        label = labels[number - 1] if number - 1 < len(labels) else str(number)
        if number > 1:
            doc.add("")
        suffix = "" if label == str(number) else f" (printed {label})"
        page_lines.append(doc.add_anchor(f"=== PAGE {number}{suffix} ===", number))
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = ""
            doc.warn(f"page {number}: text extraction failed ({exc})")
        text = re.sub(r"\n{3,}", "\n\n", text.replace("\x00", "")).strip()
        if text:
            chars += len(text)
            doc.add(text)
        if number % 50 == 0:
            progress(f"read {number}/{total} pages")
    if chars == 0:
        raise ExtractError(
            "No text found - the PDF looks scanned. Run OCR first (e.g. `brew install ocrmypdf && ocrmypdf in.pdf out.pdf`) "
            "and ingest the OCR'd file.",
            4,
        )
    if chars / max(total, 1) < 50:
        doc.warn(
            "Very little text per page - parts of this PDF may be scanned images (consider OCR with ocrmypdf)."
        )

    def page_span(page: int) -> tuple[int, int]:
        return page_lines[page - 1], (
            page_lines[page] - 1 if page < total else len(doc.lines)
        )

    entries: list[tuple[int, str, int]] = []

    def walk(items, level: int) -> None:
        for item in items:
            if isinstance(item, list):
                walk(item, level + 1)
                continue
            try:
                index = reader.get_destination_page_number(item)
            except Exception:
                continue
            title = str(getattr(item, "title", "") or "").strip()
            if title and index is not None and 0 <= index < total:
                entries.append((level, title, index + 1))

    try:
        walk(reader.outline, 1)
    except Exception as exc:
        doc.warn(f"Could not read PDF bookmarks ({exc})")

    if entries:
        doc.detected = "bookmarks"
        for level, title, page in entries:
            first, last = page_span(page)
            doc.mark(
                level,
                title,
                find_heading_line(doc.lines, title, first + 1, last) or first,
            )
    else:
        for title, line in heuristic_chapters(doc, page_lines, page_span):
            doc.mark(1, title, line)
        if doc.marks:
            doc.detected = "heuristic"
    return doc


def heuristic_chapters(
    doc: Doc, page_lines: list[int], page_span
) -> list[tuple[str, int]]:
    found: dict[str, list[tuple[str, int]]] = {"chapter": [], "part": []}
    last_number = {"chapter": 0, "part": 0}
    for page in range(1, len(page_lines) + 1):
        first, last = page_span(page)
        body = [
            (n, doc.lines[n - 1].strip())
            for n in range(first + 1, last + 1)
            if doc.lines[n - 1].strip()
        ]
        if sum(1 for _, text in body if CHAPTER_LINE_RE.match(text)) >= 3:
            continue  # table-of-contents page
        head = body[:6]
        for position, (line_no, text) in enumerate(head):
            match = CHAPTER_LINE_RE.match(text)
            if not match:
                continue
            kind, raw, rest = (
                match.group(1).lower(),
                match.group(2),
                match.group(3).strip(),
            )
            if rest and re.search(r"\d\s*$", rest):
                continue  # contents entry such as "Chapter 2 Basics ..... 31"
            number = int(raw) if raw.isdigit() else roman_to_int(raw)
            if number <= last_number[kind]:
                continue  # running header or repeated mention
            if (
                not rest
                and position + 1 < len(head)
                and len(head[position + 1][1]) <= 100
            ):
                rest = head[position + 1][1]
            found[kind].append(
                (f"{kind.title()} {raw}" + (f": {rest}" if rest else ""), line_no)
            )
            last_number[kind] = number
            break
    for kind in ("chapter", "part"):
        if len(found[kind]) >= 2:
            return found[kind]
    return []


def pptx_shape_lines(shape) -> list[str]:
    if hasattr(shape, "shapes"):  # group shape
        return [line for child in shape.shapes for line in pptx_shape_lines(child)]
    if getattr(shape, "has_table", False):
        return [
            " | ".join(" ".join(cell.text.split()) for cell in row.cells)
            for row in shape.table.rows
        ]
    if getattr(shape, "has_text_frame", False):
        lines = []
        for paragraph in shape.text_frame.paragraphs:
            text = " ".join(paragraph.text.split())
            if text:
                lines.append("  " * paragraph.level + "- " + text)
        return lines
    return []


def pptx_sections(prs, slides) -> list[tuple[str, int]]:
    namespace = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
    index = {slide.slide_id: number for number, slide in enumerate(slides, 1)}
    groups = []
    try:
        for section in prs.part._element.iter(f"{namespace}section"):
            numbers = sorted(
                index[int(e.get("id"))]
                for e in section.iter(f"{namespace}sldId")
                if int(e.get("id")) in index
            )
            if numbers:
                groups.append(
                    (section.get("name") or f"Section {len(groups) + 1}", numbers[0])
                )
    except Exception:
        return []
    return sorted(groups, key=lambda group: group[1]) if len(groups) > 1 else []


def extract_pptx(path: Path, per_group: int) -> Doc:
    pptx = need("pptx", "python-pptx", "PowerPoint files")
    try:
        prs = pptx.Presentation(str(path))
    except Exception as exc:
        raise ExtractError(f"Could not open the presentation: {exc}") from exc
    slides = list(prs.slides)
    if not slides:
        raise ExtractError("The presentation has no slides.", 4)
    titles: list[str] = []
    slide_lines: list[int] = []
    doc = Doc("", "pptx", "slide")
    for number, slide in enumerate(slides, 1):
        title_shape = slide.shapes.title
        title = (
            " ".join(title_shape.text_frame.text.split())
            if title_shape is not None and title_shape.has_text_frame
            else ""
        )
        titles.append(title)
        if number > 1:
            doc.add("")
        slide_lines.append(
            doc.add_anchor(
                (
                    f"=== SLIDE {number}: {title} ==="
                    if title
                    else f"=== SLIDE {number} ==="
                ),
                number,
            )
        )
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            for line in pptx_shape_lines(shape):
                doc.add(line)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                doc.add("Speaker notes: " + notes)
    doc.title = (
        (prs.core_properties.title or "").strip()
        or next((t for t in titles if t), "")
        or pretty(path.stem)
    )

    groups = pptx_sections(prs, slides)
    detected = "sections"
    if not groups:
        starts = [
            n
            for n, slide in enumerate(slides, 1)
            if "section" in (slide.slide_layout.name or "").lower()
        ]
        if starts:
            groups = [(titles[n - 1] or f"Section from slide {n}", n) for n in starts]
            if starts[0] != 1:
                groups.insert(0, (titles[0] or "Opening", 1))
            detected = "section-slides"
    if len(groups) < 2:
        size = max(1, per_group)
        groups = []
        for start in range(1, len(slides) + 1, size):
            end = min(start + size - 1, len(slides))
            label = next((t for t in titles[start - 1 : end] if t), "")
            groups.append(
                (f"Slides {start}-{end}" + (f": {label}" if label else ""), start)
            )
        detected = "slide-groups"
    for title, start in groups:
        doc.mark(1, title, slide_lines[start - 1])
    doc.detected = detected
    return doc


def extract_docx(path: Path) -> Doc:
    docx = need("docx", "python-docx", "Word files")
    from docx.table import Table

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractError(f"Could not open the Word document: {exc}") from exc
    title = (document.core_properties.title or "").strip()
    doc = Doc(title or pretty(path.stem), "docx", "heading")
    blocks = (
        document.iter_inner_content()
        if hasattr(document, "iter_inner_content")
        else document.paragraphs
    )
    for block in blocks:
        if isinstance(block, Table):
            for row in block.rows:
                doc.add(" | ".join(" ".join(cell.text.split()) for cell in row.cells))
            doc.add("")
            continue
        text = " ".join(block.text.split())
        if not text:
            continue
        try:
            style = (block.style.name if block.style is not None else "") or ""
        except Exception:
            style = ""
        heading = re.match(r"Heading (\d)", style)
        if style == "Title":
            doc.title = title or text
            doc.add(f"# {text}")
        elif heading:
            level = int(heading.group(1))
            doc.add("")
            doc.mark(level, text, doc.add("#" * min(level + 1, 6) + " " + text))
        elif "List" in style:
            doc.add(f"- {text}")
        else:
            doc.add(text)
    if not any(line.strip() for line in doc.lines):
        raise ExtractError("No text found in the Word document.", 4)
    if doc.marks:
        doc.detected = "headings"
    return doc


class _HTMLToText(HTMLParser):
    SKIP = {
        "script",
        "style",
        "noscript",
        "svg",
        "nav",
        "footer",
        "aside",
        "form",
        "template",
    }
    BLOCK = {
        "p",
        "div",
        "section",
        "article",
        "main",
        "br",
        "tr",
        "blockquote",
        "table",
        "ul",
        "ol",
        "dl",
        "dt",
        "dd",
        "figure",
        "figcaption",
        "hr",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._pre = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag == "pre":
            self._pre += 1
            self.parts.append("\n```\n")
        elif re.fullmatch(r"h[1-6]", tag):
            self.parts.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag == "title":
            self._in_title = False
        elif tag == "pre":
            self._pre = max(0, self._pre - 1)
            self.parts.append("\n```\n")
        elif re.fullmatch(r"h[1-6]", tag) or tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._skip:
            return
        if self._in_title:
            self.title += data
        elif self._pre:
            self.parts.append(data)
        else:
            self.parts.append(re.sub(r"\s+", " ", data))


def extract_html(path: Path) -> Doc:
    parser = _HTMLToText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    cleaned, fenced = [], False
    for line in "".join(parser.parts).splitlines():
        if line.strip() == "```":
            fenced = not fenced
        cleaned.append(line.rstrip() if fenced else line.strip())
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(cleaned)).strip()
    if not text:
        raise ExtractError("No text found in the HTML file.", 4)
    doc = Doc(" ".join(parser.title.split()) or pretty(path.stem), "html", "heading")
    doc.add(text)
    add_markdown_marks(doc)
    return doc


def text_doc(path: Path) -> Doc:
    raw = path.read_text(encoding="utf-8", errors="replace")
    doc = Doc(
        pretty(path.stem),
        "text" if path.suffix.lower() == ".txt" else "markdown",
        "heading",
    )
    doc.lines = re.split(r"\r\n|\r|\n", raw)
    if doc.lines and doc.lines[-1] == "":
        doc.lines.pop()
    add_markdown_marks(doc)
    return doc


# ---------------------------------------------------------------- transcripts


def to_seconds(stamp: str) -> float:
    parts = stamp.replace(",", ".").split(":")
    hours = int(parts[-3]) if len(parts) == 3 else 0
    return hours * 3600 + int(parts[-2]) * 60 + float(parts[-1])


def parse_subtitles(raw: str) -> list[tuple[float, float, str]]:
    cues: list[tuple[float, float, str]] = []
    start = end = None
    buffer: list[str] = []
    last_line = None
    for line in raw.splitlines():
        match = CUE_RE.match(line)
        if match:
            if start is not None and buffer:
                cues.append((start, end, " ".join(buffer)))
            start, end, buffer = (
                to_seconds(match.group(1)),
                to_seconds(match.group(2)),
                [],
            )
            continue
        if start is None:
            continue
        text = " ".join(unescape(TAG_RE.sub("", line)).split())
        if not text:
            if buffer:
                cues.append((start, end, " ".join(buffer)))
            start, buffer = None, []
            continue
        if text == last_line:
            continue  # auto-captions repeat the previous line in every cue
        buffer.append(text)
        last_line = text
    if start is not None and buffer:
        cues.append((start, end, " ".join(buffer)))
    return cues


def parse_json3(raw: str) -> list[tuple[float, float, str]]:
    cues = []
    for event in json.loads(raw).get("events", []):
        text = " ".join(
            "".join(seg.get("utf8", "") for seg in event.get("segs") or []).split()
        )
        if text:
            start = event.get("tStartMs", 0) / 1000
            cues.append((start, start + event.get("dDurationMs", 0) / 1000, text))
    return cues


def parse_subtitle_file(path: Path) -> list[tuple[float, float, str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    return parse_json3(raw) if path.suffix.lower() == ".json3" else parse_subtitles(raw)


def group_cues(cues, target: float = 30.0) -> list[tuple[float, str]]:
    paragraphs, current, current_start = [], [], None
    for start, _end, text in cues:
        if current_start is None:
            current_start = start
        current.append(text)
        elapsed = start - current_start
        if (
            elapsed >= target and text.endswith((".", "?", "!"))
        ) or elapsed >= target * 2:
            paragraphs.append((current_start, " ".join(current)))
            current, current_start = [], None
    if current:
        paragraphs.append((current_start, " ".join(current)))
    return paragraphs


def transcript_doc(
    title, cues, duration, chapters, segment_minutes: float, source_label: str
) -> Doc:
    if not cues:
        raise ExtractError("No speech or captions were found in this media.", 4)
    cues = sorted(cues)
    doc = Doc(title, "transcript", "time")
    doc.duration = duration or max(cue[1] for cue in cues)
    if chapters:
        chapters = sorted(chapters)
        doc.detected = "chapters"
    else:
        span = max(60.0, segment_minutes * 60)
        chapters = [
            (i * span, "") for i in range(max(1, math.ceil(doc.duration / span)))
        ]
        doc.detected = "segments"
    doc.add(f"# {title}")
    doc.add(f"Source: {source_label} | Duration: {hms(doc.duration)}")
    starts = [start for start, _ in chapters]
    buckets: list[list] = [[] for _ in chapters]
    for cue in cues:
        buckets[max(0, bisect.bisect_right(starts, cue[0]) - 1)].append(cue)
    kept = [
        (
            start,
            name,
            bucket,
            chapters[i + 1][0] if i + 1 < len(chapters) else doc.duration,
        )
        for i, ((start, name), bucket) in enumerate(zip(chapters, buckets))
        if bucket
    ]  # drop silent stretches
    for number, (start, name, bucket, end) in enumerate(kept, 1):
        name = name or f"Part {number}"
        label = (
            f"CHAPTER {number}: {name}"
            if doc.detected == "chapters"
            else f"PART {number}"
        )
        doc.add("")
        doc.mark(
            1,
            name,
            doc.add_anchor(f"=== {label} [{hms(start)}] ===", start),
            start=start,
            end=end,
        )
        for paragraph_start, text in group_cues(bucket):
            doc.add_anchor(
                f"[{hms(paragraph_start)}] {text}", paragraph_start, marker=False
            )
    return doc


def transcribe(media: Path | None, cache: Path, args) -> tuple[list, float]:
    if cache.exists() and not args.force:
        data = json.loads(cache.read_text(encoding="utf-8"))
        return [tuple(cue) for cue in data["cues"]], data.get("duration")
    if media is None:
        raise ExtractError("Nothing to transcribe.")
    whisper = need(
        "faster_whisper", "faster-whisper", "transcribing audio/video without captions"
    )
    progress(
        f"Transcribing {media.name} with Whisper '{args.whisper_model}' (the first run downloads the model)..."
    )
    try:
        model = whisper.WhisperModel(
            args.whisper_model, device="auto", compute_type="auto"
        )
    except Exception as exc:
        raise ExtractError(
            f"Could not load the Whisper model '{args.whisper_model}' ({type(exc).__name__}: {str(exc).splitlines()[0][:200]}). "
            "Models download from huggingface.co - check network/proxy access, or pass --whisper-model with a local "
            "faster-whisper model folder, or put a .srt/.vtt captions file next to the video (same name)."
        ) from exc
    try:
        segments, info = model.transcribe(
            str(media), language=args.lang, vad_filter=True
        )
        cues, reported = [], 0.0
        for segment in segments:
            text = segment.text.strip()
            if text:
                cues.append((round(segment.start, 2), round(segment.end, 2), text))
            if segment.end - reported >= 300:
                reported = segment.end
                progress(f"  transcribed {hms(segment.end)} / {hms(info.duration)}")
    except Exception as exc:
        raise ExtractError(f"Transcription failed for {media.name}: {exc}") from exc
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(
        json.dumps(
            {
                "model": args.whisper_model,
                "language": info.language,
                "duration": info.duration,
                "cues": cues,
            }
        ),
        encoding="utf-8",
    )
    return cues, info.duration


def find_sidecar_subtitles(media: Path) -> Path | None:
    candidates = sorted(
        p
        for p in media.parent.iterdir()
        if p.is_file()
        and p.suffix.lower() in SUB_EXT
        and (p.stem == media.stem or p.name.startswith(media.stem + "."))
    )
    return candidates[0] if candidates else None


def extract_media(path: Path, out: Path, name: str, args) -> Doc:
    sidecar = find_sidecar_subtitles(path)
    if sidecar:
        progress(f"Using captions from {sidecar.name}")
        cues, duration = parse_subtitle_file(sidecar), None
    else:
        cues, duration = transcribe(path, out / f"{name}.whisper.json", args)
    return transcript_doc(
        pretty(path.stem), cues, duration, None, args.segment_minutes, path.name
    )


def youtube_video(yt_dlp, url: str, out: Path, args) -> Doc:
    lang = args.lang or "en"
    options = {
        **YDL_BASE,
        "skip_download": True,
        "writesubtitles": True,
        "writeautomaticsub": True,
        "subtitleslangs": [lang],
        "subtitlesformat": "json3/vtt/srt/best",
        "ignoreerrors": True,
        "outtmpl": str(out / "%(id)s.%(ext)s"),
    }
    try:
        with yt_dlp.YoutubeDL(options) as ydl:
            info = ydl.extract_info(url, download=True)
        if (
            info is None
        ):  # ignoreerrors swallowed an extraction error; retry strictly to surface it
            with yt_dlp.YoutubeDL({**YDL_BASE, "skip_download": True}) as ydl:
                info = ydl.extract_info(url, download=False)
    except Exception as exc:
        raise ExtractError(f"yt-dlp could not read {url}: {exc}") from exc
    video_id = safe_name(info.get("id") or url)
    subtitles = sorted(
        p
        for p in out.iterdir()
        if p.name.startswith(f"{video_id}.")
        and p.suffix.lower() in {".json3", ".vtt", ".srt"}
    )
    preferred = [p for p in subtitles if p.name.split(".")[-2] == lang] or subtitles
    duration = info.get("duration")
    if preferred:
        cues = parse_subtitle_file(preferred[0])
    else:
        cache = out / f"{video_id}.whisper.json"
        audio = None
        if args.force or not cache.exists():
            progress(
                "No captions available - downloading audio for local transcription..."
            )
            audio_options = {
                **YDL_BASE,
                "format": "bestaudio/best",
                "outtmpl": str(out / "%(id)s.audio.%(ext)s"),
            }
            try:
                with yt_dlp.YoutubeDL(audio_options) as ydl:
                    ydl.extract_info(url, download=True)
            except Exception as exc:
                raise ExtractError(f"Audio download failed: {exc}") from exc
            audio = next(
                (p for p in out.iterdir() if p.name.startswith(f"{video_id}.audio.")),
                None,
            )
        cues, duration = transcribe(audio, cache, args)
        if audio is not None:
            audio.unlink(missing_ok=True)
    chapters = [
        (float(c.get("start_time") or 0), c.get("title") or f"Chapter {i}")
        for i, c in enumerate(info.get("chapters") or [], 1)
    ]
    doc = transcript_doc(
        info.get("title") or video_id,
        cues,
        duration,
        chapters,
        args.segment_minutes,
        url,
    )
    doc.meta_id = video_id
    return doc


# ---------------------------------------------------------------- orchestration


def write_text(doc: Doc, out: Path, name: str) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    target = out / f"{name}.txt"
    target.write_text("\n".join(doc.lines) + "\n", encoding="utf-8")
    return target


def write_outline(outline: dict, out: Path, name: str) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{name}.outline.json"
    path.write_text(
        json.dumps(outline, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    summary = {key: value for key, value in outline.items() if key != "structure"}
    summary["outline"] = display_path(path)
    summary["topLevel"] = [
        f"{unit['id']} {unit['title']} ({unit.get('loc') or 'no loc'})"
        for unit in outline["structure"][:40]
    ]
    return summary


def finish(
    doc: Doc, text_path: Path, source_label: str, out: Path, name: str, args
) -> dict:
    structure, skipped = build_structure(doc, args.keep_all)
    return write_outline(
        {
            "generatedBy": GENERATOR,
            "title": doc.title,
            "source": source_label,
            "sourceText": display_path(text_path),
            "kind": doc.kind,
            "locator": doc.locator,
            "detected": doc.detected if structure else "none",
            "totalUnits": count_units(structure),
            "totalLines": len(doc.lines),
            "skipped": skipped,
            "warnings": doc.warnings,
            "structure": structure,
        },
        out,
        name,
    )


def finish_collection(
    title: str, source_label: str, items, out: Path, name: str, args
) -> dict:
    structure, warnings = [], []
    for doc, text_path, chapter_title in items:
        units, _ = build_structure(doc, args.keep_all)
        leaves = (
            units[0]["sections"] if len(units) == 1 and units[0]["sections"] else units
        )
        number = len(structure) + 1
        structure.append(
            {
                "id": f"ch-{number}",
                "title": chapter_title,
                "source": display_path(text_path),
                "loc": doc.span_loc(),
                "lines": [1, len(doc.lines)],
                "sections": [
                    {
                        "id": f"{number}.{k}",
                        "title": leaf["title"],
                        "loc": leaf["loc"],
                        "lines": leaf["lines"],
                    }
                    for k, leaf in enumerate(leaves, 1)
                ],
            }
        )
        warnings += [f"{chapter_title}: {warning}" for warning in doc.warnings]
    return write_outline(
        {
            "generatedBy": GENERATOR,
            "title": title,
            "source": source_label,
            "sourceText": None,
            "kind": "collection",
            "locator": "mixed",
            "detected": "files",
            "totalUnits": count_units(structure),
            "skipped": [],
            "warnings": warnings[:20],
            "structure": structure,
        },
        out,
        name,
    )


def extract_file(path: Path, out: Path, args, name: str) -> tuple[Doc, Path]:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXT:
        return text_doc(path), path
    if suffix == ".pdf":
        doc = extract_pdf(path)
    elif suffix == ".pptx":
        doc = extract_pptx(path, args.slides_per_group)
    elif suffix == ".docx":
        doc = extract_docx(path)
    elif suffix in HTML_EXT:
        doc = extract_html(path)
    elif suffix in SUB_EXT:
        doc = transcript_doc(
            pretty(path.stem),
            parse_subtitle_file(path),
            None,
            None,
            args.segment_minutes,
            path.name,
        )
    elif suffix in MEDIA_EXT:
        doc = extract_media(path, out, name, args)
    elif suffix in CONVERT_HINTS:
        raise ExtractError(
            f"{path.name}: {suffix} is not supported directly. {CONVERT_HINTS[suffix]}"
        )
    else:
        raise ExtractError(f"{path.name}: unsupported file type '{suffix}'.")
    return doc, write_text(doc, out, name)


def unzip_safely(archive: Path, destination: Path, force: bool) -> Path:
    root = destination.resolve()
    with zipfile.ZipFile(archive) as bundle:
        members = []
        for member in bundle.infolist():
            if member.is_dir() or member.filename.startswith("__MACOSX/"):
                continue
            target = (root / member.filename).resolve()
            if not target.is_relative_to(root):
                raise ExtractError(
                    f"Refusing to extract unsafe path from zip: {member.filename}"
                )
            members.append((member, target))
        for member, target in members:
            if target.exists() and not force:
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with bundle.open(member) as source, open(target, "wb") as sink:
                shutil.copyfileobj(source, sink)
    return destination


def ingest_folder(folder: Path, out: Path, args) -> dict:
    files = sorted(
        (
            p
            for p in folder.rglob("*")
            if p.is_file()
            and p.suffix.lower() in SUPPORTED
            and not p.name.startswith(".")
            and not {"_extracted", "__MACOSX"} & set(p.relative_to(folder).parts)
        ),
        key=natural_key,
    )
    media = {(p.parent, p.stem) for p in files if p.suffix.lower() in MEDIA_EXT}
    files = [
        p
        for p in files  # captions next to a video are used by that video, not ingested twice
        if not (
            p.suffix.lower() in SUB_EXT
            and {(p.parent, p.stem), (p.parent, p.name.split(".")[0])} & media
        )
    ]
    if not files:
        raise ExtractError(f"No supported files found in {folder}")
    items = []
    for number, path in enumerate(files, 1):
        progress(f"[{number}/{len(files)}] {path.relative_to(folder)}")
        try:
            doc, text_path = extract_file(
                path, out, args, slugify(path.relative_to(folder).as_posix())
            )
        except ExtractError as exc:
            progress(f"  skipped: {exc}")
            continue
        items.append((doc, text_path, pretty(path.stem)))
    if not items:
        raise ExtractError("None of the files in the folder could be extracted.", 4)
    return finish_collection(
        pretty(folder.name),
        display_path(folder),
        items,
        out,
        safe_name(args.name) if args.name else slugify(folder.name),
        args,
    )


def ingest_youtube(url: str, out: Path, args) -> dict:
    yt_dlp = need("yt_dlp", "yt-dlp", "YouTube links")
    if re.search(r"youtube\.com/playlist\?", url, re.I):
        try:
            with yt_dlp.YoutubeDL(
                {**YDL_BASE, "noplaylist": False, "extract_flat": "in_playlist"}
            ) as ydl:
                info = ydl.extract_info(url, download=False)
        except Exception as exc:
            raise ExtractError(f"yt-dlp could not read the playlist: {exc}") from exc
        entries = [
            entry for entry in info.get("entries") or [] if entry and entry.get("id")
        ]
        if not entries:
            raise ExtractError("The playlist is empty or private.", 4)
        items = []
        for number, entry in enumerate(entries, 1):
            progress(f"[{number}/{len(entries)}] {entry.get('title') or entry['id']}")
            doc = youtube_video(
                yt_dlp, f"https://www.youtube.com/watch?v={entry['id']}", out, args
            )
            items.append((doc, write_text(doc, out, safe_name(entry["id"])), doc.title))
        title = info.get("title") or "YouTube playlist"
        return finish_collection(
            title,
            url,
            items,
            out,
            safe_name(args.name) if args.name else slugify(title),
            args,
        )
    doc = youtube_video(yt_dlp, url, out, args)
    name = safe_name(args.name or doc.meta_id)
    return finish(doc, write_text(doc, out, name), url, out, name, args)


def run(args) -> dict:
    source = args.source.strip()
    if re.match(r"^https?://", source, re.I):
        if not YOUTUBE_RE.match(source):
            raise ExtractError(
                "Only YouTube URLs are handled here. Save web articles as markdown with the fetch tool instead."
            )
        if not args.out_dir:
            raise ExtractError(
                "--out-dir is required for URLs (e.g. learning-materials/<slug>/_extracted)."
            )
        return ingest_youtube(source, Path(args.out_dir).expanduser(), args)
    path = Path(source).expanduser()
    if not path.exists():
        raise ExtractError(f"Not found: {path}")
    if path.is_file() and path.suffix.lower() == ".zip":
        path = unzip_safely(path, path.parent / path.stem, args.force)
    if path.is_dir():
        out = Path(args.out_dir).expanduser() if args.out_dir else path / "_extracted"
        return ingest_folder(path, out, args)
    out = (
        Path(args.out_dir).expanduser() if args.out_dir else path.parent / "_extracted"
    )
    name = safe_name(args.name) if args.name else slugify(path.stem)
    doc, text_path = extract_file(path, out, args, name)
    return finish(doc, text_path, display_path(path), out, name, args)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("source", help="file, folder, .zip, or YouTube URL")
    parser.add_argument(
        "--out-dir", help="output folder (default: <source folder>/_extracted)"
    )
    parser.add_argument("--name", help="base name for the output files")
    parser.add_argument(
        "--segment-minutes",
        type=float,
        default=10,
        help="video segment length when there are no chapters",
    )
    parser.add_argument(
        "--slides-per-group",
        type=int,
        default=10,
        help="slides per unit when a deck has no sections",
    )
    parser.add_argument(
        "--whisper-model",
        default="small",
        help="faster-whisper model name (tiny, base, small, medium, large-v3) or a local model folder",
    )
    parser.add_argument(
        "--lang",
        help="caption/transcription language code, e.g. en (default: en for YouTube, auto-detect for Whisper)",
    )
    parser.add_argument(
        "--keep-all",
        action="store_true",
        help="keep cover/contents/index entries in the outline",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="redo cached work (zip extraction, transcription)",
    )
    args = parser.parse_args(argv)
    try:
        summary = run(args)
    except MissingDependency as exc:
        print(
            f'MISSING DEPENDENCY: {exc}\nInstall it with: "{sys.executable}" -m pip install {exc.package}',
            file=sys.stderr,
        )
        return 3
    except ExtractError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return exc.code
    except KeyboardInterrupt:
        print("Interrupted.", file=sys.stderr)
        return 130
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
